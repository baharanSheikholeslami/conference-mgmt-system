"""
سندهای مصنوعی برای سناریوهای سرتاسری (فاز ۴) و ارزیابی عامل اعتبارسنجی (فاز ۵).

هر سند یک «فرم پیشنهاد پروژه» به قالب DOCX است که از روی چک‌لیست agents/checklist.py ساخته
می‌شود. با variant می‌شود عمداً بخشی را حذف یا فیلدی را خالی گذاشت؛ چون می‌دانیم چه چیزی حذف
شده، پاسخ درست عامل اعتبارسنجی از پیش معلوم است (expected_missing).
"""
import io
from dataclasses import dataclass

from django.core.files.uploadedfile import SimpleUploadedFile

from agents import checklist


@dataclass(frozen=True)
class Topic:
    """موضوع یک پیشنهاد مصنوعی: عنوان‌ها و چند جمله برای هر بخش."""
    title_fa: str
    title_en: str
    intro: str
    related: str
    approach: str
    evaluation: str


TOPICS = {
    "llm_agents": Topic(
        "دستیار هوشمند پاسخ‌گویی به آیین‌نامه‌های آموزشی با مدل‌های زبانی بزرگ",
        "An LLM-based Assistant for Answering Academic Regulation Questions",
        "دانشجویان برای یافتن پاسخ پرسش‌های آیین‌نامه‌ای باید اسناد طولانی را بخوانند یا با کارشناس "
        "آموزش مکاتبه کنند. این پروژه یک دستیار گفتگومحور مبتنی بر مدل زبانی بزرگ و بازیابی-افزوده "
        "تولید می‌سازد که با استناد به بند مرجع پاسخ می‌دهد و پردازش زبان طبیعی فارسی را به کار می‌گیرد.",
        "سامانه‌های پرسش و پاسخ دانشگاهی معمولاً مبتنی بر قواعد یا جست‌وجوی کلیدواژه‌ای هستند. "
        "پژوهش‌های اخیر نشان داده‌اند که بازیابی-افزوده تولید خطای مدل زبانی را کاهش می‌دهد.",
        "اسناد آیین‌نامه قطعه‌بندی و با یک مدل embedding چندزبانه در پایگاه دانش برداری نمایه می‌شوند. "
        "پرسش کاربر به بردار تبدیل، نزدیک‌ترین بندها بازیابی و به مدل زبانی داده می‌شود.",
        "یک مجموعه‌ی پنجاه‌تایی پرسش و پاسخ آزمایشی تهیه و دقت پاسخ‌ها و درستی استناد به بند مرجع "
        "به‌صورت انسانی سنجیده می‌شود."),
    "databases": Topic(
        "طراحی پایگاه‌داده‌ی توزیع‌شده برای سامانه‌ی ثبت‌نام دانشگاه",
        "Designing a Distributed Database for a University Registration System",
        "سامانه‌ی ثبت‌نام دانشگاه در روزهای انتخاب واحد با بار بسیار زیاد روبه‌روست. این پروژه یک "
        "طراحی پایگاه‌داده‌ی توزیع‌شده با تکرار داده و افراز افقی جدول‌ها ارائه می‌کند و مهندسی "
        "نرم‌افزار لایه‌ی دسترسی به داده را بهبود می‌دهد.",
        "پایگاه‌داده‌های رابطه‌ای توزیع‌شده و راهکارهای افراز داده در ادبیات بررسی شده‌اند.",
        "طرح‌واره‌ی داده بازطراحی، کلید افراز انتخاب و تراکنش‌های توزیع‌شده با پروتکل دومرحله‌ای پیاده می‌شوند.",
        "زمان پاسخ و توان عملیاتی سامانه با شبیه‌سازی بار روز انتخاب واحد اندازه‌گیری می‌شود."),
}

# variant ← کلیدهای چک‌لیستی که عمداً حذف یا خالی می‌شوند
VARIANTS = {
    "complete": (),
    "no_evaluation": ("evaluation",),
    "no_references": ("references",),
    "no_schedule_resources": ("schedule", "resources"),
    "empty_student_fields": ("student_number", "student_email"),
}

_SECTION_HEADINGS = {
    "introduction": "۱. مقدمه",
    "related_work": "۲. پروژه‌ها و سامانه‌های مشابه",
    "approach": "۳. روش انجام پروژه",
    "evaluation": "۴. روش ارزیابی",
    "schedule": "۵. مراحل انجام و زمان‌بندی پروژه",
    "resources": "۶. امکانات لازم",
    "references": "۷. مراجع و منابع",
}


def expected_missing(variant):
    """پاسخ درست عامل اعتبارسنجی برای این variant: کلیدهایی که باید «ناقص» اعلام شوند."""
    return set(VARIANTS[variant])


def _body(topic, key):
    return {
        "introduction": topic.intro,
        "related_work": topic.related,
        "approach": topic.approach,
        "evaluation": topic.evaluation,
        "schedule": "مرحله‌ی اول: مطالعه‌ی پیشینه (دو هفته). مرحله‌ی دوم: طراحی و پیاده‌سازی (شش هفته). "
                    "مرحله‌ی سوم: آزمایش و ارزیابی (دو هفته). مرحله‌ی چهارم: نگارش گزارش (یک هفته).",
        "resources": "یک رایانه‌ی شخصی، زبان برنامه‌نویسی پایتون، پایگاه‌داده‌ی PostgreSQL و ابزارهای متن‌باز.",
        "references": "[1] Y. Gao et al., \"Retrieval-Augmented Generation for Large Language Models: "
                      "A Survey,\" arXiv:2312.10997, 2023.\n[2] W. W. Wan Ishak, \"A Comparative Review of "
                      "Conference Management System,\" JTOM, 2023.",
    }[key]


def proposal_docx(topic="llm_agents", variant="complete", student="دانشجوی آزمایشی",
                  student_number="۴۰۰۳۹۹۹۹", advisor="استاد راهنمای آزمایشی"):
    """یک فرم پیشنهاد پروژه به قالب DOCX؛ خروجی آماده‌ی بارگذاری (SimpleUploadedFile)."""
    from docx import Document

    t = TOPICS[topic]
    drop = set(VARIANTS[variant])
    fields = [
        ("student_name", "نام و نام خانوادگی", student),
        ("student_number", "شماره دانشجویی", student_number),
        ("student_email", "رایانامه (ایمیل) دانشجو", "student@aut.ac.ir"),
        ("term", "نیم‌سال و سال تحصیلی ثبت‌نام پروژه", "نیم‌سال اول از سال تحصیلی ۱۴۰۵"),
    ]
    document = Document()
    document.add_paragraph("فرم پیشنهاد پروژه کارشناسی — دانشکده مهندسی کامپیوتر")
    document.add_paragraph("مشخصات دانشجو")
    grid = document.add_table(rows=len(fields), cols=2)
    for row, (key, label, value) in zip(grid.rows, fields):
        row.cells[0].text, row.cells[1].text = label, ("" if key in drop else value)
    document.add_paragraph("استاد راهنمای پروژه")
    document.add_paragraph(f"نام و نام خانوادگی: {'' if 'advisor_name' in drop else advisor}")
    if "title" not in drop:
        document.add_paragraph(f"عنوان فارسی: {t.title_fa}")
        document.add_paragraph(f"عنوان انگلیسی: {t.title_en}")
    document.add_paragraph("تعریف پروژه")
    for key in [k for k, *_ in checklist.SECTIONS if k != "title"]:
        if key in drop:
            continue
        document.add_paragraph(_SECTION_HEADINGS[key])
        for line in _body(t, key).split("\n"):
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return SimpleUploadedFile(f"proposal-{topic}-{variant}.docx", buffer.getvalue(),
                             content_type="application/vnd.openxmlformats-officedocument."
                                          "wordprocessingml.document")


def thesis_docx(topic="llm_agents"):
    """یک پایان‌نامه‌ی کوتاه نمونه (فقط برای بارگذاری؛ عامل اعتبارسنجی پایان‌نامه را بررسی نمی‌کند)."""
    from docx import Document

    t = TOPICS[topic]
    document = Document()
    for line in (t.title_fa, "فصل ۱: مقدمه", t.intro, "فصل ۲: روش", t.approach, "فصل ۳: ارزیابی",
                 t.evaluation):
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return SimpleUploadedFile(f"thesis-{topic}.docx", buffer.getvalue())
