"""
قراردادن تصویر امضای الکترونیک روی سند تأییدشده.

نسخه‌ی اصلی سند هرگز تغییر نمی‌کند؛ یک «نسخه‌ی امضاشده» کنار امضا (Signature.signed_file) ساخته
می‌شود. اگر همان نسخه قبلاً امضا شده باشد (مثلاً استاد راهنما و بعد داور)، امضای تازه روی آخرین
نسخه‌ی امضاشده می‌نشیند تا همه‌ی امضاها در یک فایل باشند.
  • DOCX: نام، نقش، تاریخ و تصویر امضا زیر «محل امضا»ی قالب (یا اگر نبود، در انتهای سند) می‌نشیند
  • PDF : تصویر امضا و تاریخ در پایین صفحه‌ی آخر چاپ می‌شود
اگر ساخت نسخه‌ی امضاشده ممکن نباشد (فایل خراب یا قالب دیگر)، امضا ثبت می‌ماند و فقط فایل ساخته نمی‌شود.
"""
import io
import logging
from pathlib import Path

from django.core.files.base import ContentFile
from django.utils import timezone

from .jalali import to_jalali

log = logging.getLogger("workflow")

SIGNATURE_WIDTH_CM = 4


def _read(field):
    field.open("rb")
    try:
        return field.read()
    finally:
        field.close()


def _base(record):
    """سندی که امضا روی آن می‌نشیند: آخرین نسخه‌ی امضاشده‌ی همین نسخه، یا خود فایل اصلی."""
    previous = (record.version.signatures.exclude(pk=record.pk).exclude(signed_file="")
                .order_by("-created_at").first())
    return _read(previous.signed_file if previous else record.version.file)


def _caption(record):
    return f"{record.get_role_display()}: {record.signer} — {to_jalali(timezone.localdate())}"


# محل امضای هر نقش در قالب نمونه (workflow/static/workflow/templates)
PLACEHOLDERS = {"advisor": "امضای استاد راهنما", "referee": "امضای داور"}
HINT = "(امضای الکترونیک"


def _placeholder(document, role):
    """پاراگراف «محل امضا»ی این نقش در سند، یا None."""
    label = PLACEHOLDERS.get(role)
    return next((p for p in document.paragraphs if label and label in p.text), None)


def _sign_docx(data, image, record):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    document = Document(io.BytesIO(data))
    caption = document.add_paragraph(_caption(record))
    picture = document.add_paragraph()
    picture.add_run().add_picture(io.BytesIO(image), width=Cm(SIGNATURE_WIDTH_CM))
    for p in (caption, picture):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    anchor = _placeholder(document, record.role)
    if anchor is not None:
        # امضا درست زیر «محل امضا»ی قالب می‌نشیند و راهنمای «(امضای الکترونیک … این‌جا)» برداشته می‌شود
        following = anchor._p.getnext()
        if following is not None and following.xpath("string(.)").strip().startswith(HINT):
            following.getparent().remove(following)
        anchor._p.addnext(picture._p)
        anchor._p.addnext(caption._p)
    elif not record.version.signatures.exclude(pk=record.pk).exclude(signed_file="").exists():
        heading = document.add_paragraph()
        heading.add_run("امضای الکترونیک").bold = True
        heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        heading._p.addnext(caption._p)
        caption._p.addnext(picture._p)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _sign_pdf(data, image, record):
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    reader = PdfReader(io.BytesIO(data))
    writer = PdfWriter(clone_from=reader)
    last = writer.pages[-1]
    width, height = float(last.mediabox.width), float(last.mediabox.height)
    # هر امضای بعدی کمی بالاتر از قبلی تا روی هم نیفتند
    slot = record.version.signatures.exclude(pk=record.pk).exclude(signed_file="").count()
    overlay = io.BytesIO()
    pdf = canvas.Canvas(overlay, pagesize=(width, height))
    picture = ImageReader(io.BytesIO(image))
    w, h = picture.getSize()
    draw_w = SIGNATURE_WIDTH_CM * cm
    draw_h = draw_w * h / w
    x, y = 2 * cm, 2 * cm + slot * (draw_h + 1.2 * cm)
    pdf.drawImage(picture, x, y + .5 * cm, width=draw_w, height=draw_h, mask="auto")
    pdf.setFont("Helvetica", 8)
    pdf.drawString(x, y, f"e-signed {timezone.localdate().isoformat()} #{record.signer_id}")
    pdf.save()
    last.merge_page(PdfReader(io.BytesIO(overlay.getvalue())).pages[0])
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


SIGNERS = {".docx": _sign_docx, ".pdf": _sign_pdf}


def sign_version(record):
    """نسخه‌ی امضاشده‌ی سند را می‌سازد و در record.signed_file ذخیره می‌کند. خروجی: موفق بود؟"""
    extension = Path(record.version.file.name).suffix.lower()
    signer = SIGNERS.get(extension)
    if signer is None:
        return False
    try:
        signed = signer(_base(record), _read(record.image), record)
    except Exception as e:             # فایل خراب، تصویر نامعتبر، … — امضا ثبت می‌ماند
        log.warning("signing version %s failed: %s", record.version_id, e)
        return False
    name = f"case{record.version.case_id}-{record.version.kind}-v{record.version.number}-signed{extension}"
    record.signed_file.save(name, ContentFile(signed), save=True)
    return True
