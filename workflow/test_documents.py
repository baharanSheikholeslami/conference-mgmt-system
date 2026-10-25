"""آزمون‌های صفحه‌ی پرونده‌ی تازه: باکس اسناد، نمایشگر سند، نظر روی بخش‌ها، تأیید با امضا و بازگرداندن."""
import io
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from . import services
from .enums import DocKind, ReviewDecision, State
from .models import DocumentComment, Signature
from .test_portal import PortalTestCase
from .test_services import pdf, signature_png
from .views import CONTENT_TYPES

SAMPLES = Path(__file__).resolve().parent.parent / "docs" / "samples"


def sample_docx(kind="proposal"):
    return SimpleUploadedFile(f"{kind}.docx", (SAMPLES / f"{kind}-sara-nouri.docx").read_bytes())


class DocumentReviewTests(PortalTestCase):
    def advisor_review(self, file=None):
        case = self.choose_advisor(self.new_case())
        services.upload_document(case, self.student, DocKind.PROPOSAL, file or sample_docx())
        return services.perform(case, "submit_proposal", self.student)

    def panel(self, user, case, version, **params):
        self.login(user)
        return self.client.get(reverse("case_detail", args=[case.pk]), {"doc": version.pk, **params})

    def test_word_file_opens_in_its_own_format_with_a_section_picker(self):
        case = self.advisor_review()
        version = case.versions.get()
        page = self.panel(self.advisor, case, version)
        inline = reverse("document_inline", args=[version.pk])
        self.assertContains(page, f'data-src="{inline}"')               # فایل Word با docx-preview نمایش داده می‌شود
        self.assertContains(page, "docx-preview.min.js")
        self.assertIn((5, "۴. روش ارزیابی"), page.context["panel"]["section_choices"])
        self.assertContains(page, "تأیید و ثبت امضای الکترونیک")
        response = self.client.get(inline)
        self.assertEqual(response["Content-Type"], CONTENT_TYPES["docx"])
        self.assertFalse(response.has_header("Content-Disposition") and "attachment" in response["Content-Disposition"])
        page = self.panel(self.student, case, version)                 # دانشجو فرم نظر ندارد
        self.assertNotContains(page, 'name="section"')
        self.assertIsNone(page.context["panel"]["decision"])

    def test_pdf_file_opens_in_a_frame(self):
        case = self.advisor_review(pdf())
        version = case.versions.get()
        page = self.panel(self.advisor, case, version)
        self.assertContains(page, f'<iframe class="doc-frame" src="{reverse("document_inline", args=[version.pk])}"')
        self.assertNotContains(page, "docx-preview.min.js")

    def test_text_view_splits_the_document_into_sections_with_comment_forms(self):
        case = self.advisor_review()
        version = case.versions.get()
        page = self.panel(self.advisor, case, version, view="text")
        sections = page.context["panel"]["sections"]
        self.assertIn("۴. روش ارزیابی", [s.title for s in sections])
        self.assertContains(page, "نظر روی این بخش")
        page = self.panel(self.student, case, version, view="text")
        self.assertNotContains(page, "نظر روی این بخش")

    def test_unreadable_file_still_opens_with_a_general_comment_box(self):
        case = self.advisor_review(pdf())
        page = self.panel(self.advisor, case, case.versions.get())
        self.assertIsNone(page.context["panel"]["sections"])
        self.assertTrue(page.context["panel"]["show_file"])
        self.assertEqual(page.context["panel"]["section_choices"], [(0, "کل سند")])
        self.assertContains(page, 'name="section"')

    def test_comment_on_a_section_is_shown_to_the_student(self):
        case = self.advisor_review()
        version = case.versions.get()
        self.login(self.advisor)
        url = reverse("document_comment", args=[version.pk])
        response = self.client.post(url, {"text": "معیار ارزیابی را کمّی کنید", "section": 5,
                                          "section_title": "۴. روش ارزیابی", "view": "text"})
        self.assertRedirects(response, f"{reverse('case_detail', args=[case.pk])}?doc={version.pk}&view=text#s-5",
                             fetch_redirect_response=False)
        self.assertContains(self.panel(self.student, case, version), "معیار ارزیابی را کمّی کنید")
        # از نمای «فایل» فقط شماره‌ی بخش می‌آید؛ سرفصل از متن سند خوانده می‌شود
        self.login(self.advisor)
        response = self.client.post(url, {"text": "مرجع کم است", "section": 8, "view": "file"})
        self.assertRedirects(response, f"{reverse('case_detail', args=[case.pk])}?doc={version.pk}#comments",
                             fetch_redirect_response=False)
        self.assertEqual(DocumentComment.objects.get(text="مرجع کم است").section_title, "۷. مراجع و منابع")
        self.login(self.student)                                       # دانشجو نظر نمی‌گذارد
        self.assertEqual(self.client.post(url, {"text": "x"}).status_code, 403)
        self.login(self.outsider)
        self.assertEqual(self.client.post(url, {"text": "x"}).status_code, 404)

    def test_approve_requires_a_signature_image_and_places_it_on_the_document(self):
        case = self.advisor_review()
        version = case.versions.get()
        self.login(self.advisor)
        url = reverse("document_decide", args=[version.pk])
        self.client.post(url, {"decision": "approve"})                 # بدون امضا
        case.refresh_from_db()
        self.assertEqual(case.state, State.ADVISOR_REVIEW)
        self.client.post(url, {"decision": "approve", "signature": pdf("sign.pdf")})
        case.refresh_from_db()
        self.assertEqual(case.state, State.ADVISOR_REVIEW)

        response = self.client.post(url, {"decision": "approve", "signature": signature_png()})
        # پس از تأیید، نمایشگر با نسخه‌ی امضاشده باز می‌شود
        self.assertRedirects(response, f"{reverse('case_detail', args=[case.pk])}?doc={version.pk}",
                             fetch_redirect_response=False)
        case.refresh_from_db()
        self.assertEqual(case.state, State.EDU_REGISTRATION)
        signature = Signature.objects.get()
        self.assertEqual((signature.signer, signature.role), (self.advisor, "advisor"))
        from docx import Document
        signed = Document(signature.signed_file.path)
        self.assertEqual(len(signed.inline_shapes), 1)
        self.assertIn(str(self.advisor), "\n".join(p.text for p in signed.paragraphs))
        for user, status in ((self.student, 200), (self.edu, 200), (self.outsider, 404)):
            self.login(user)
            self.assertEqual(self.client.get(reverse("signature_image", args=[signature.pk])).status_code,
                             status, user.username)
            self.assertEqual(self.client.get(reverse("signed_download", args=[signature.pk])).status_code,
                             status, user.username)

    def test_signed_copy_is_shown_and_downloadable_after_approval(self):
        case = self.advisor_review()
        version = case.versions.get()
        inline = reverse("document_inline", args=[version.pk])
        self.login(self.student)
        self.assertEqual(self.client.get(inline, {"signed": 1}).status_code, 404)     # هنوز امضا نشده
        page = self.client.get(reverse("case_detail", args=[case.pk]))
        self.assertNotContains(page, "دانلود نسخه‌ی امضاشده")

        self.login(self.advisor)
        self.client.post(reverse("document_decide", args=[version.pk]),
                         {"decision": "approve", "signature": signature_png()})
        signature = Signature.objects.get()
        self.login(self.student)
        page = self.client.get(reverse("case_detail", args=[case.pk]))
        self.assertContains(page, reverse("signed_download", args=[signature.pk]))     # در باکس سند
        self.assertContains(page, "دانلود نسخه‌ی امضاشده")
        page = self.panel(self.student, case, version)
        self.assertContains(page, f'data-src="{inline}?signed=1"')                     # نمایشگر نسخه‌ی امضاشده را نشان می‌دهد
        self.assertContains(page, "دانلود امضاشده")
        page = self.panel(self.student, case, version, copy="original")
        self.assertContains(page, f'data-src="{inline}"')
        response = self.client.get(inline, {"signed": 1})
        from docx import Document
        self.assertEqual(len(Document(io.BytesIO(b"".join(response.streaming_content))).inline_shapes), 1)
        response = self.client.get(reverse("signed_download", args=[signature.pk]))
        self.assertIn(f"case{case.pk}-proposal-v1-signed.docx", response["Content-Disposition"])
        self.login(self.outsider)
        self.assertEqual(self.client.get(inline, {"signed": 1}).status_code, 404)

    def test_plain_approve_form_redirects_to_the_signature_panel(self):
        case = self.advisor_review()
        self.login(self.advisor)
        response = self.client.get(reverse("case_action", args=[case.pk, "advisor_approve"]))
        self.assertRedirects(response, f"{reverse('case_detail', args=[case.pk])}"
                                       f"?doc={case.versions.get().pk}&panel=sign")

    def test_return_uses_section_comments_when_no_note_is_written(self):
        case = self.advisor_review()
        version = case.versions.get()
        url = reverse("document_decide", args=[version.pk])
        self.login(self.advisor)
        self.client.post(url, {"decision": "revise"})                  # نه توضیح، نه نظر
        case.refresh_from_db()
        self.assertEqual(case.state, State.ADVISOR_REVIEW)
        services.add_comment(case, self.advisor, version, "منابع ناقص است", 8, "۷. مراجع و منابع")
        self.client.post(url, {"decision": "revise"})
        case.refresh_from_db()
        self.assertEqual(case.state, State.DRAFTING)
        self.assertIn("منابع ناقص است", services.pending_feedback(case)["items"][0]["note"])

    def test_referee_votes_from_the_viewer(self):
        case = self.to_referee_review()
        version = case.versions.get()
        self.login(self.ref1)
        self.client.post(reverse("document_decide", args=[version.pk]),
                         {"decision": "approve", "signature": signature_png()})
        case.refresh_from_db()
        self.assertEqual(case.state, State.COUNCIL_APPROVAL)
        self.assertEqual(version.reviews.get().decision, ReviewDecision.APPROVE)
        self.assertEqual(Signature.objects.get().role, "referee")
        self.client.post(reverse("document_decide", args=[version.pk]), {"decision": "revise", "note": "x"})
        self.assertEqual(version.reviews.count(), 1)                    # دیگر تصمیمی با او نیست

    def test_old_versions_get_no_decision_or_comments(self):
        case = self.advisor_review()
        case = services.perform(case, "advisor_return", self.advisor, note="اصلاح")
        old = case.versions.get()
        services.upload_document(case, self.student, DocKind.PROPOSAL, sample_docx())
        case = services.perform(case, "submit_proposal", self.student)
        self.assertIsNone(services.doc_decision(case, self.advisor, old))
        self.assertFalse(services.can_comment(case, self.advisor, old))
        with self.assertRaises(services.NotAllowed):
            services.add_comment(case, self.advisor, old, "x")
        self.assertFalse(DocumentComment.objects.exists())


class DocumentBoxTests(PortalTestCase):
    def test_empty_boxes_offer_the_template_and_upload(self):
        case = self.choose_advisor(self.new_case())
        self.login(self.student)
        page = self.client.get(reverse("case_detail", args=[case.pk]))
        self.assertContains(page, "proposal-template.docx")
        self.assertContains(page, "thesis-template.docx")
        self.assertContains(page, reverse("document_upload", args=[case.pk]))
        boxes = page.context["boxes"]
        self.assertEqual([b["can_upload"] for b in boxes], [True, False])

    def test_after_upload_the_box_offers_submit_and_reupload(self):
        case = self.choose_advisor(self.new_case())
        services.upload_document(case, self.student, DocKind.PROPOSAL, pdf())
        self.login(self.student)
        page = self.client.get(reverse("case_detail", args=[case.pk]))
        self.assertContains(page, reverse("case_action", args=[case.pk, "submit_proposal"]))
        self.assertContains(page, "بارگذاری نسخه‌ی تازه")

    def test_template_files_exist(self):
        from django.contrib.staticfiles import finders
        from .views import TEMPLATES
        for path in TEMPLATES.values():
            self.assertTrue(finders.find(path), path)


class DashboardLayoutTests(PortalTestCase):
    def test_student_without_advisor_is_offered_the_choice(self):
        case = self.new_case()
        self.login(self.student)
        page = self.client.get(reverse("dashboard"))
        self.assertTemplateUsed(page, "workflow/dashboard_student.html")
        self.assertContains(page, reverse("case_action", args=[case.pk, "select_advisor"]))

    def test_student_without_a_project_sees_the_empty_state(self):
        self.login(self.student)
        page = self.client.get(reverse("dashboard"))
        self.assertContains(page, "هنوز پروژه‌ای ندارید")
        self.assertContains(page, reverse("case_create"))

    def test_student_reminders_are_only_proposal_thesis_and_defense(self):
        self.to_referee_review()
        self.login(self.student)
        kinds = {t["deadline"].kind for t in self.client.get(reverse("dashboard")).context["timers"]}
        self.assertTrue(kinds)
        self.assertLessEqual(kinds, {"proposal_submission", "thesis_delivery"})

    def test_each_role_gets_its_dashboard(self):
        self.to_referee_review()
        for user, template in ((self.advisor, "dashboard_faculty"), (self.head, "dashboard_faculty"),
                               (self.edu, "dashboard_edu"), (self.external, "dashboard_student")):
            self.login(user)
            self.assertTemplateUsed(self.client.get(reverse("dashboard")), f"workflow/{template}.html")
        self.login(self.head)
        page = self.client.get(reverse("dashboard"))
        self.assertContains(page, "همه‌ی پروژه‌های گروه")
        self.login(self.edu)
        page = self.client.get(reverse("dashboard"))
        self.assertEqual(sum(p["count"] for p in page.context["phase_rows"]), 1)
