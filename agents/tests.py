"""
آزمون‌های لایه‌ی عامل‌ها (سمت پورتال). n8n و مدل زبانی این‌جا «ساختگی» هستند: تنها نقطه‌ی
تماس پورتال با n8n تابع agents.n8n.post است که با mock جایگزین می‌شود؛ پس این آزمون‌ها بدون
n8n و بدون اینترنت اجرا می‌شوند و قرارداد دو طرف (چه می‌فرستیم، چه می‌پذیریم) را می‌سنجند.
"""
import io
import json
import logging
from datetime import timedelta
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from workflow import rules
from workflow import services as wf
from workflow.enums import DeadlineKind, DocKind, State, UserType
from workflow.forms import AssignRefereesForm
from workflow.models import User
from workflow.statemachine import TRANSITIONS
from workflow.test_services import MEDIA, WorkflowTestCase, pdf

from . import checklist, chunking, extract, n8n, services
from .models import (AssistantMessage, DocumentValidation, Notification, PolicyChunk,
                     SuggestionRun)

TOKEN = "test-token-0123456789abcdef"
STARTED = {"message": "Workflow was started"}
BODY = ("این پروژه یک سامانه‌ی چندعامله برای مدیریت گردش‌کار پروژه‌ی کارشناسی طراحی و "
        "پیاده‌سازی می‌کند و آن را با سناریوهای آزمایشی ارزیابی می‌کند. ") * 2
PROPOSAL = ["عنوان فارسی: سامانه هوشمند مدیریت کنفرانس", "عنوان انگلیسی: Intelligent Conference System",
            "۱. مقدمه", BODY, "۲. پروژه‌ها و سامانه‌های مشابه", BODY, "۳. روش انجام پروژه", BODY,
            "۴. روش ارزیابی", BODY, "۵. مراحل انجام و زمان‌بندی پروژه", BODY,
            "۶. امکانات لازم", BODY, "۷. مراجع و منابع", "[1] Y. Gao et al., RAG Survey, 2023."]
FORM_TABLE = [("نام و نام خانوادگی", "بهاران شیخ‌الاسلامی"), ("شماره دانشجویی", "۴۰۰۳۱۰۸۹"),
              ("رایانامه (ایمیل) دانشجو", "baharan.sh@aut.ac.ir")]


def setUpModule():
    logging.disable(logging.CRITICAL)        # هشدارهای مورد انتظار (n8n خاموش، PDF خراب) خروجی آزمون را شلوغ نکنند


def tearDownModule():
    logging.disable(logging.NOTSET)


def docx(paragraphs=PROPOSAL, table=FORM_TABLE, name="proposal.docx"):
    from docx import Document
    document = Document()
    if table:
        grid = document.add_table(rows=len(table), cols=2)
        for row, (label, value) in zip(grid.rows, table):
            row.cells[0].text, row.cells[1].text = label, value
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def text_pdf(text, name="proposal.pdf"):
    """یک PDF یک‌صفحه‌ای واقعی با متن لاتین (برای آزمودن مسیر pypdf)."""
    stream = b"BT /F1 11 Tf 40 780 Td (" + text.encode("latin-1") + b") Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (i, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return SimpleUploadedFile(name, out, content_type="application/pdf")


# ───────────────────────── توابع خالص ─────────────────────────

class ExtractTests(SimpleTestCase):
    def test_normalize_unifies_persian_text(self):
        self.assertEqual(extract.normalize("كتاب  علي\u200f  ـــ\n\n\n\nمي‌رود"), "کتاب علی\n\nمی‌رود")
        self.assertEqual(extract.normalize("ﻻ ﷲ"), "لا الله")           # شکل‌های نمایشی PDF ← حرف پایه

    def test_docx_text_includes_tables_in_document_order(self):
        text = extract.extract_text(docx())
        self.assertIn("شماره دانشجویی | ۴۰۰۳۱۰۸۹", text)
        self.assertLess(text.index("شماره دانشجویی"), text.index("۱. مقدمه"))
        self.assertLess(text.index("۱. مقدمه"), text.index("۷. مراجع و منابع"))

    def test_pdf_text(self):
        self.assertIn("Hello proposal", extract.extract_text(text_pdf("Hello proposal")))

    def test_pdf_right_to_left_repairs(self):
        self.assertEqual(extract._repair_rtl("شماره دانشجویی۹۸۰۱۳۰۰۴ ، تاریخ ۱/۹/۴۰۴۱ ، بند ۲-۵ ، نمره ۵٫۸۱"),
                         "شماره دانشجویی۴۰۰۳۱۰۸۹ ، تاریخ ۱۴۰۴/۹/۱ ، بند ۵-۲ ، نمره ۱۸٫۵")
        self.assertEqual(extract._repair_rtl("رایانامه )ایمیل( دانشجو\nf(x) = (a)"),
                         "رایانامه (ایمیل) دانشجو\nf(x) = (a)")

    def test_unsupported_or_broken_files_raise_a_persian_error(self):
        for file in (SimpleUploadedFile("old.doc", b"\xd0\xcf\x11\xe0"),
                     SimpleUploadedFile("broken.pdf", b"not a pdf"),
                     SimpleUploadedFile("broken.docx", b"not a zip")):
            with self.assertRaises(extract.ExtractionError, msg=file.name):
                extract.extract_text(file)


class ChunkingTests(SimpleTestCase):
    def test_splits_on_headings_and_keeps_the_reference(self):
        text = ("# رویه‌ی آزمایشی\n\n۱- هدف\nهدف این رویه تعیین مراحل پروژه است.\n"
                "۵-۲ دانشجو موظف است حداکثر دو ماه پس از ثبت‌نام پروژه، فرم تکمیل‌شده‌ی پیشنهاد را "
                "به آموزش دانشکده تحویل دهد.\nتبصره ۱: این مهلت تمدید نمی‌شود مگر با موافقت شورای آموزشی.\n"
                "۱. این یک مورد از فهرستی بلند است که نباید به‌تنهایی عنوان حساب شود چون یک جمله‌ی کامل و طولانی است.")
        chunks = chunking.split(text)
        self.assertEqual([c.ref for c in chunks], ["۱- هدف", "بند ۵-۲", "تبصره ۱"])
        self.assertTrue(chunks[1].text.startswith("۵-۲ دانشجو موظف است"))
        self.assertIn("یک مورد از فهرستی بلند", chunks[2].text)       # مورد فهرست زیر همان تبصره می‌ماند

    def test_long_sections_are_wrapped_at_sentence_boundaries(self):
        sentence = "این یک جمله‌ی آزمایشی درباره‌ی مهلت‌های پروژه‌ی کارشناسی است."
        chunks = chunking.split("## بند بلند\n" + " ".join([sentence] * 40))
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(c.text) <= chunking.MAX_CHARS and c.ref == "بند بلند" for c in chunks))
        self.assertTrue(all(c.text.endswith(".") for c in chunks))

    def test_bundled_knowledge_file_gives_one_chunk_per_note(self):
        from agents.management.commands.index_policy import KNOWLEDGE_DIR
        chunks = chunking.split(extract.extract_text(KNOWLEDGE_DIR / "FORM-PROPOSAL.md"))
        refs = [c.ref for c in chunks]
        self.assertEqual(len(refs), 8)
        self.assertTrue(refs[3].startswith("توضیح ۴"))
        self.assertIn("حداقل سه ماه", chunks[3].text)


class ReminderRuleTests(SimpleTestCase):
    def test_stage_for_each_number_of_days_left(self):
        expected = {30: None, 15: None, 14: "d14", 8: "d14", 7: "d7", 4: "d7", 3: "d3", 2: "d3",
                    1: "d1", 0: "d0", -1: "overdue1", -7: "overdue1", -8: "overdue2", -15: "overdue3"}
        self.assertEqual({d: rules.reminder_stage(d) for d in expected}, expected)


# ───────────────────────── پایه‌ی آزمون‌های پایگاه‌داده ─────────────────────────

@override_settings(AGENTS_ENABLED=True, AGENT_TOKEN=TOKEN, MEDIA_ROOT=MEDIA,
                   PORTAL_PUBLIC_URL="http://portal.test",
                   PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AgentTestCase(WorkflowTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        for user in User.objects.all():
            user.email = f"{user.username}@example.com"
            user.first_name = user.username
            user.save()
        for user in (cls.student, cls.external, cls.advisor, cls.ref1, cls.ref2, cls.head, cls.outsider, cls.edu):
            user.refresh_from_db()

    def setUp(self):
        patcher = mock.patch("agents.n8n.post", return_value=STARTED)
        self.n8n = patcher.start()
        self.addCleanup(patcher.stop)

    def calls(self, path):
        """بدنه‌ی همه‌ی درخواست‌هایی که به این Webhook فرستاده شده است."""
        return [c.args[1] for c in self.n8n.call_args_list if c.args[0] == path]

    def api(self, name, data=None, args=(), token=TOKEN, method="post", query=""):
        url = reverse(name, args=args) + query
        headers = {"X-Agent-Token": token} if token else {}
        if method == "get":
            return self.client.get(url, headers=headers)
        return self.client.post(url, json.dumps(data or {}), content_type="application/json",
                                headers=headers)

    def drafting_case(self):
        return self.do(self.new_case(), "select_advisor", self.student, advisor=self.advisor)

    def upload(self, case, file=None):
        """بارگذاری پیشنهاد، همراه با اجرای قلاب‌های پس از commit (مثل یک درخواست واقعی)."""
        with self.captureOnCommitCallbacks(execute=True):
            return wf.upload_document(case, case.author, DocKind.PROPOSAL, file or docx())

    def act(self, case, action, actor, **kwargs):
        with self.captureOnCommitCallbacks(execute=True):
            return self.do(case, action, actor, **kwargs)


class ApiSecurityTests(AgentTestCase):
    def test_every_agent_endpoint_requires_the_shared_token(self):
        endpoints = [("api_validation_result", (1,)), ("api_referee_suggestions_result", (1,)),
                     ("api_rag_search", ()), ("api_notification_record", ())]
        for name, args in endpoints:
            for token in ("", "wrong-token"):
                self.assertEqual(self.api(name, args=args, token=token).status_code, 403, name)
        self.assertEqual(self.api("api_reminders_due", method="get", token="").status_code, 403)

    @override_settings(AGENT_TOKEN="")
    def test_an_unset_token_locks_the_api_instead_of_opening_it(self):
        self.assertEqual(self.api("api_rag_search", token="").status_code, 403)

    def test_wrong_method_and_bad_json_are_rejected(self):
        self.assertEqual(self.api("api_rag_search", method="get").status_code, 405)
        response = self.client.post(reverse("api_rag_search"), "[1, 2]", content_type="application/json",
                                    headers={"X-Agent-Token": TOKEN})
        self.assertEqual(response.status_code, 400)


# ───────────────────────── ۱. اعتبارسنجی سند ─────────────────────────

class ValidationTests(AgentTestCase):
    def found(self, *skip):
        return [{"key": i["key"], "evidence": f"شاهد {i['label']}"} for i in checklist.ITEMS
                if i["key"] not in skip]

    def test_upload_sends_text_and_checklist_to_n8n(self):
        version = self.upload(self.drafting_case())
        (payload,) = self.calls("validate-proposal")
        validation = DocumentValidation.objects.get(version=version)
        self.assertEqual((validation.status, payload["validation_id"]), ("pending", validation.pk))
        self.assertIn("شماره دانشجویی | ۴۰۰۳۱۰۸۹", payload["document"]["text"])
        self.assertEqual([i["key"] for i in payload["checklist"]], [i["key"] for i in checklist.ITEMS])
        self.assertEqual(len([i for i in payload["checklist"] if i["group"] == "section"]), 8)

    def test_thesis_upload_is_not_validated(self):
        case = self.to_in_progress()
        with self.captureOnCommitCallbacks(execute=True):
            wf.upload_document(case, self.student, DocKind.THESIS, pdf("thesis.pdf"))
        self.assertEqual(self.calls("validate-proposal"), [])

    def test_result_with_missing_items_is_reported_as_issues(self):
        version = self.upload(self.drafting_case())
        validation = version.validation
        response = self.api("api_validation_result", args=(validation.pk,), data={
            "found": self.found("evaluation", "student_email", "references") + [{"key": "invented", "evidence": "x"}],
            "unverified": [{"key": "references", "evidence": "ساختگی"}],
            "windows": 3, "failed_windows": 0, "model": "test-model"})
        self.assertEqual(response.json(), {"status": "issues"})
        validation.refresh_from_db()
        self.assertEqual([i["key"] for i in validation.problems], ["evaluation", "references", "student_email"])
        self.assertEqual(len(validation.items), len(checklist.ITEMS))          # کلید ساختگی مدل نادیده گرفته شد
        notes = {i["key"]: i["note"] for i in validation.items}
        self.assertIn("تأیید نشد", notes["references"])
        self.assertIn("شاهد در متن", notes["introduction"])
        self.assertEqual((validation.model, validation.error), ("test-model", ""))
        self.assertIsNotNone(validation.finished_at)

    def test_complete_document_passes_and_partial_failure_is_flagged(self):
        validation = self.upload(self.drafting_case()).validation
        self.api("api_validation_result", args=(validation.pk,),
                 data={"found": self.found(), "windows": 4, "failed_windows": 1})
        validation.refresh_from_db()
        self.assertEqual(validation.status, "passed")
        self.assertIn("ناقص", validation.error)

    def test_agent_error_and_unknown_id(self):
        validation = self.upload(self.drafting_case()).validation
        self.api("api_validation_result", args=(validation.pk,), data={"error": "مدل پاسخ نداد"})
        validation.refresh_from_db()
        self.assertEqual((validation.status, validation.error), ("error", "مدل پاسخ نداد"))
        self.assertEqual(self.api("api_validation_result", args=(99999,)).status_code, 404)

    def test_file_without_readable_text_never_reaches_n8n(self):
        case = self.drafting_case()
        for file in (pdf(), docx(paragraphs=["خیلی کوتاه"], table=None), text_pdf("short")):
            validation = self.upload(case, file).validation
            self.assertEqual(validation.status, "error", file.name)
            self.assertTrue(validation.error)
        self.assertEqual(self.calls("validate-proposal"), [])

    def test_n8n_being_down_never_breaks_the_upload(self):
        self.n8n.side_effect = n8n.Unavailable("connection refused")
        case = self.drafting_case()
        version = self.upload(case)
        self.assertEqual(version.validation.status, "error")
        self.assertEqual(self.do(case, "submit_proposal", self.student).state, State.ADVISOR_REVIEW)

    def test_validation_is_advice_only_and_never_blocks_submission(self):
        case = self.drafting_case()
        validation = self.upload(case).validation
        self.api("api_validation_result", args=(validation.pk,), data={"found": []})
        self.assertEqual(DocumentValidation.objects.get(pk=validation.pk).status, "issues")
        self.assertEqual(self.do(case, "submit_proposal", self.student).state, State.ADVISOR_REVIEW)

    def test_case_page_shows_feedback_and_retry_button(self):
        case = self.drafting_case()
        validation = self.upload(case).validation
        self.client.force_login(self.student)
        page = self.client.get(reverse("case_detail", args=[case.pk]))
        self.assertContains(page, "data-agent-waiting")
        self.api("api_validation_result", args=(validation.pk,), data={"found": self.found("schedule")})
        page = self.client.get(reverse("case_detail", args=[case.pk]))
        self.assertContains(page, "۱ مورد نیازمند توجه")
        self.assertNotContains(page, "data-agent-waiting")
        self.assertContains(page, reverse("validation_retry", args=[validation.version_id]))

    def test_pending_result_goes_stale(self):
        validation = self.upload(self.drafting_case()).validation
        validation.requested_at = timezone.now() - timedelta(minutes=11)
        self.assertTrue(validation.is_stale)
        self.assertFalse(validation.is_waiting)

    def test_retry_is_limited_to_author_and_edu(self):
        case = self.drafting_case()
        version = self.upload(case)
        url = reverse("validation_retry", args=[version.pk])
        for user, status in ((self.advisor, 403), (self.outsider, 404), (self.edu, 302), (self.student, 302)):
            self.client.force_login(user)
            self.assertEqual(self.client.post(url).status_code, status, user.username)
        self.assertEqual(len(self.calls("validate-proposal")), 3)
        self.assertEqual(self.client.get(url).status_code, 405)


# ───────────────────────── ۲. دستیار RAG ─────────────────────────

class RagTests(AgentTestCase):
    MODEL = "test-embed"

    def chunk(self, position, ref, vector, model=MODEL, code="AUT-PR-3210"):
        return PolicyChunk.objects.create(doc_code=code, position=position, ref=ref, text=f"متن {ref}",
                                          embedding=vector, embedding_model=model)

    def test_index_document_chunks_embeds_and_replaces(self):
        from agents.management.commands.index_policy import KNOWLEDGE_DIR
        self.n8n.side_effect = lambda path, payload, timeout=None: {
            "model": self.MODEL, "vectors": [[1.0, 0.0, float(len(t) % 7)] for t in payload["texts"]]}
        self.chunk(1, "قدیمی", [1, 1, 1], code="FORM-PROPOSAL")
        count, model = services.index_document(KNOWLEDGE_DIR / "FORM-PROPOSAL.md", "FORM-PROPOSAL", "فرم")
        self.assertEqual((count, model), (8, self.MODEL))
        chunks = PolicyChunk.objects.filter(doc_code="FORM-PROPOSAL")
        self.assertEqual([c.position for c in chunks], list(range(1, 9)))
        self.assertFalse(chunks.filter(ref="قدیمی").exists())
        (payload,) = self.calls("rag-embed")
        self.assertTrue(payload["texts"][0].startswith("توضیح ۱"))             # عنوان بند هم بردارسازی می‌شود

    def test_index_fails_cleanly_when_embedding_is_incomplete(self):
        from agents.management.commands.index_policy import KNOWLEDGE_DIR
        self.n8n.return_value = {"model": self.MODEL, "vectors": [[1, 0, 0]]}
        with self.assertRaises(services.AgentError):
            services.index_document(KNOWLEDGE_DIR / "FORM-PROPOSAL.md", "FORM-PROPOSAL")
        self.assertFalse(PolicyChunk.objects.exists())

    def test_search_returns_nearest_chunks_by_cosine_similarity(self):
        self.chunk(1, "بند دور", [0, 1, 0])
        near = self.chunk(2, "بند نزدیک", [1, 0.1, 0])
        self.chunk(3, "بند میانه", [1, 1, 0])
        self.chunk(4, "مدل دیگر", [1, 0, 0, 0], model="other-model")           # ابعاد و مدل متفاوت
        response = self.api("api_rag_search", data={"embedding": [1, 0, 0], "model": self.MODEL, "k": 2})
        chunks = response.json()["chunks"]
        self.assertEqual([c["ref"] for c in chunks], ["بند نزدیک", "بند میانه"])
        self.assertEqual(chunks[0]["id"], near.pk)
        self.assertAlmostEqual(chunks[0]["score"], 0.995, places=3)
        self.assertAlmostEqual(chunks[1]["score"], 0.7071, places=3)

    def test_search_rejects_bad_vectors_and_unindexed_models(self):
        self.chunk(1, "بند", [1, 0, 0])
        for data in ({"embedding": [], "model": self.MODEL}, {"embedding": ["a"], "model": self.MODEL},
                     {"embedding": [1, 0, 0], "model": "never-indexed"}):
            response = self.api("api_rag_search", data=data)
            self.assertEqual(response.status_code, 409, data)
            self.assertIn("error", response.json())

    def test_chat_box_returns_answer_with_real_source_text_and_logs_it(self):
        chunk = self.chunk(1, "توضیح ۴", [1, 0, 0])
        self.n8n.return_value = {"answer": "حداقل سه ماه. [1]", "model": "test-llm", "sources": [
            {"n": 1, "chunk_id": chunk.pk, "score": 0.9, "cited": True},
            {"n": 2, "chunk_id": 987654, "score": 0.1, "cited": False}]}
        self.client.force_login(self.student)
        response = self.client.post(reverse("assistant_ask"), json.dumps({"question": "  فاصله تا دفاع؟ "}),
                                    content_type="application/json")
        data = response.json()
        self.assertEqual((data["ok"], data["answer"]), (True, "حداقل سه ماه. [1]"))
        self.assertEqual(data["sources"], [{"n": 1, "doc": "AUT-PR-3210", "ref": "توضیح ۴",
                                            "text": "متن توضیح ۴", "score": 0.9, "cited": True}])
        self.assertEqual(self.calls("rag-ask"), [{"question": "فاصله تا دفاع؟"}])
        message = AssistantMessage.objects.get()
        self.assertEqual((message.user, message.ok, message.model), (self.student, True, "test-llm"))

    def test_chat_box_errors_are_friendly_and_logged(self):
        self.client.force_login(self.student)

        def ask(question):
            return self.client.post(reverse("assistant_ask"), json.dumps({"question": question}),
                                    content_type="application/json").json()

        self.assertIn("نمایه نشده", ask("سلام")["error"])                         # پایگاه دانش خالی
        self.chunk(1, "بند", [1, 0, 0])
        self.assertFalse(ask("   ")["ok"])
        self.assertFalse(ask("ب" * 501)["ok"])
        self.assertEqual(self.calls("rag-ask"), [])
        self.n8n.side_effect = n8n.Unavailable("timeout")
        self.assertIn("در دسترس نیست", ask("مهلت تحویل پیشنهاد؟")["error"])
        self.assertFalse(AssistantMessage.objects.get().ok)

    def test_chat_box_needs_login(self):
        response = self.client.post(reverse("assistant_ask"), "{}", content_type="application/json")
        self.assertEqual(response.status_code, 302)
        self.client.force_login(self.student)
        self.assertEqual(self.client.get(reverse("assistant_ask")).status_code, 405)
        self.assertContains(self.client.get(reverse("dashboard")), 'id="assistant-panel"')


# ───────────────────────── ۳. پیشنهاد داور ─────────────────────────

class RefereeSuggestionTests(AgentTestCase):
    def to_assignment(self):
        case = self.drafting_case()
        self.upload(case)
        case = self.do(case, "submit_proposal", self.student)
        case = self.do(case, "advisor_approve", self.advisor)
        return self.act(case, "edu_register", self.edu)

    def test_entering_referee_assignment_asks_n8n_to_rank_eligible_referees(self):
        User.objects.filter(pk=self.ref1.pk).update(expertise="پایگاه‌داده و مهندسی نرم‌افزار")
        case = self.to_assignment()
        (payload,) = self.calls("suggest-referees")
        run = SuggestionRun.objects.get(case=case)
        self.assertEqual((payload["run_id"], run.status), (run.pk, "pending"))
        self.assertEqual({c["id"] for c in payload["candidates"]}, {self.ref1.pk, self.ref2.pk, self.head.pk})
        self.assertIn({"id": self.ref1.pk, "name": "ref1", "expertise": "پایگاه‌داده و مهندسی نرم‌افزار"},
                      payload["candidates"])
        self.assertTrue(payload["query_text"].startswith("عنوان آزمایشی"))
        self.assertIn("این پروژه یک سامانه‌ی چندعامله", payload["query_text"])   # ابتدای مقدمه
        self.assertNotIn("شماره دانشجویی", payload["query_text"])               # نه صفحه‌ی اول فرم

    def test_result_is_ranked_and_filtered_by_the_portal(self):
        case = self.to_assignment()
        run = case.suggestion_run
        response = self.api("api_referee_suggestions_result", args=(run.pk,), data={"model": "m", "suggestions": [
            {"referee_id": self.head.pk, "score": None, "reason": "تخصص ثبت نشده"},
            {"referee_id": self.ref2.pk, "score": 0.41, "reason": "ارتباط کم"},
            {"referee_id": self.ref1.pk, "score": 0.83, "reason": "ارتباط زیاد"},
            {"referee_id": self.advisor.pk, "score": 0.99, "reason": "استاد راهنما واجد شرایط نیست"},
            {"referee_id": self.outsider.pk, "score": 0.95, "reason": "گروه دیگر"},
            {"referee_id": self.ref1.pk, "score": 0.10, "reason": "تکراری"},
            {"referee_id": "x", "score": 1}]})
        self.assertEqual(response.json(), {"status": "done", "saved": 3})
        rows = list(run.suggestions.values_list("referee_id", "rank", "score", "reason"))
        self.assertEqual(rows, [(self.ref1.pk, 1, 0.83, "ارتباط زیاد"), (self.ref2.pk, 2, 0.41, "ارتباط کم"),
                                (self.head.pk, 3, None, "تخصص ثبت نشده")])
        self.assertEqual(services.suggestion_ranks(case), {self.ref1.pk: 1, self.ref2.pk: 2, self.head.pk: 3})

    def test_head_sees_ranked_form_but_nothing_is_preselected(self):
        case = self.to_assignment()
        self.api("api_referee_suggestions_result", args=(case.suggestion_run.pk,), data={"suggestions": [
            {"referee_id": self.ref2.pk, "score": 0.9, "reason": "دلیل داور دو"},
            {"referee_id": self.ref1.pk, "score": 0.2, "reason": "دلیل داور یک"}]})
        form = AssignRefereesForm(case=case, transition=TRANSITIONS["assign_referees"])
        self.assertEqual(list(form.fields["referees"].queryset), [self.ref2, self.ref1, self.head])
        self.assertIn("پیشنهاد سامانه: رتبه‌ی ۱", form.fields["referees"].label_from_instance(self.ref2))
        self.assertNotIn("پیشنهاد سامانه", form.fields["referees"].label_from_instance(self.head))

        self.client.force_login(self.head)
        page = self.client.get(reverse("case_action", args=[case.pk, "assign_referees"]))
        self.assertContains(page, "دلیل داور دو")
        self.assertContains(page, "شباهت موضوعی ۹۰٪")
        self.assertContains(page, "انتخاب داور با شماست")
        self.assertNotContains(page, "checked")
        # مدیر گروه آزاد است داوری غیر از پیشنهاد اول را انتخاب کند
        self.client.post(reverse("case_action", args=[case.pk, "assign_referees"]), {"referees": [self.head.pk]})
        case.refresh_from_db()
        self.assertEqual(case.state, State.REFEREE_REVIEW)
        self.assertEqual([a.referee for a in case.referee_assignments.all()], [self.head])

    def test_agent_never_assigns_referees_itself(self):
        case = self.to_assignment()
        self.api("api_referee_suggestions_result", args=(case.suggestion_run.pk,),
                 data={"suggestions": [{"referee_id": self.ref1.pk, "score": 0.9}]})
        case.refresh_from_db()
        self.assertEqual(case.state, State.REFEREE_ASSIGNMENT)
        self.assertFalse(case.referee_assignments.exists())

    def test_error_and_refresh(self):
        case = self.to_assignment()
        run = case.suggestion_run
        self.api("api_referee_suggestions_result", args=(run.pk,), data={"error": "مدل embedding پاسخ نداد"})
        self.client.force_login(self.head)
        page = self.client.get(reverse("case_action", args=[case.pk, "assign_referees"]))
        self.assertContains(page, "مدل embedding پاسخ نداد")
        url = reverse("suggestions_refresh", args=[case.pk])
        self.assertRedirects(self.client.post(url), reverse("case_action", args=[case.pk, "assign_referees"]))
        self.assertEqual(len(self.calls("suggest-referees")), 2)
        run.refresh_from_db()
        self.assertEqual((run.status, run.error), ("pending", ""))
        for user, status in ((self.student, 403), (self.edu, 403), (self.outsider, 404)):
            self.client.force_login(user)
            self.assertEqual(self.client.post(url).status_code, status, user.username)

    def test_n8n_being_down_does_not_block_the_head(self):
        self.n8n.side_effect = n8n.Unavailable("down")
        case = self.to_assignment()
        self.assertEqual(case.suggestion_run.status, "error")
        self.assertEqual(self.do(case, "assign_referees", self.head, referees=[self.ref1]).state,
                         State.REFEREE_REVIEW)


# ───────────────────────── ۴. اطلاع‌رسانی ─────────────────────────

class NotificationTests(AgentTestCase):
    def recipients(self, items):
        return sorted(i["to"].split("@")[0] for i in items)

    def test_state_change_notifies_whoever_must_act_next(self):
        case = self.drafting_case()
        self.upload(case)
        self.n8n.reset_mock()
        case = self.act(case, "submit_proposal", self.student)
        (payload,) = self.calls("notify-event")
        (item,) = payload["notifications"]                        # فقط استاد راهنما؛ خود دانشجو اقدام‌کننده است
        event = case.events.last()
        self.assertEqual(item["key"], f"event:{event.pk}:{self.advisor.pk}")
        self.assertEqual(item["to"], "advisor@example.com")
        self.assertIn("در انتظار تأیید استاد راهنما", item["subject"])
        self.assertEqual(item["facts"]["recipient_roles"], ["استاد راهنما"])
        self.assertEqual(len(item["facts"]["your_actions"]), 2)
        self.assertIn("«ارسال پیشنهاد برای استاد راهنما» توسط student انجام شد.", item["fallback_body"])
        self.assertIn("advisor گرامی،", item["fallback_body"])
        self.assertIn(f"http://portal.test/cases/{case.pk}/", item["footer"])
        self.assertIn("پاسخ ندهید", item["footer"])
        self.assertEqual(set(item["facts"]["next_deadline"]), {"label", "date", "days_left", "status"})
        queued = Notification.objects.get()
        self.assertEqual((queued.recipient, queued.status, queued.event), (self.advisor, "queued", event))

    def test_recipients_along_the_workflow(self):
        case = self.drafting_case()
        self.upload(case)
        steps = [("submit_proposal", self.student, {}, ["advisor"]),
                 ("advisor_approve", self.advisor, {}, ["edu", "student"]),
                 ("edu_register", self.edu, {}, ["head", "student"]),
                 ("assign_referees", self.head, {"referees": [self.ref1, self.ref2]}, ["ref1", "ref2", "student"]),
                 ("submit_review", self.ref1, {"decision": "approve"}, None),      # وضعیت عوض نشد ← بدون اعلان
                 ("submit_review", self.ref2, {"decision": "revise", "note": "اصلاح شود"}, ["student"]),
                 ("cancel", self.edu, {"note": "انصراف"}, ["advisor", "student"])]   # پایان پرونده
        for action, actor, kwargs, expected in steps:
            self.n8n.reset_mock()
            case = self.act(case, action, actor, **kwargs)
            sent = self.calls("notify-event")
            if expected is None:
                self.assertEqual(sent, [], action)
            else:
                self.assertEqual(self.recipients(sent[0]["notifications"]), expected, action)

    def test_reviewer_comments_are_not_put_in_the_email(self):
        case = self.to_referee_review()
        self.n8n.reset_mock()
        self.act(case, "submit_review", self.ref1, decision="revise", note="روش ارزیابی محرمانه ضعیف است")
        (payload,) = self.calls("notify-event")
        self.assertNotIn("محرمانه", json.dumps(payload, ensure_ascii=False))

    def test_users_without_email_are_skipped(self):
        User.objects.filter(pk=self.advisor.pk).update(email="")
        case = self.drafting_case()
        self.upload(case)
        self.n8n.reset_mock()
        self.act(case, "submit_proposal", self.student)
        self.assertEqual(self.calls("notify-event"), [])
        self.assertFalse(Notification.objects.exists())

    def test_n8n_records_the_sent_email_once(self):
        case = self.drafting_case()
        self.upload(case)
        case = self.act(case, "submit_proposal", self.student)
        key = self.calls("notify-event")[0]["notifications"][0]["key"]
        for _ in range(2):                                                   # تکرار همان پاسخ بی‌ضرر است
            response = self.api("api_notification_record", data={
                "key": key, "status": "sent", "subject": "موضوع", "body": "متن پیام", "generated_by": "llm"})
            self.assertEqual(response.json()["status"], "sent")
        notification = Notification.objects.get()
        self.assertEqual((notification.status, notification.body, notification.generated_by),
                         ("sent", "متن پیام", "llm"))
        self.assertIsNotNone(notification.sent_at)
        for bad in ("", "event:1", "reminder:1:d7", "event:999999:1", "x:1:2", "event:a:b"):
            self.assertEqual(self.api("api_notification_record", data={"key": bad}).status_code, 409, bad)

    def test_failed_dispatch_is_visible_and_harmless(self):
        case = self.drafting_case()
        self.upload(case)
        self.n8n.side_effect = n8n.Unavailable("down")
        case = self.act(case, "submit_proposal", self.student)
        self.assertEqual(case.state, State.ADVISOR_REVIEW)
        self.assertEqual(Notification.objects.get().status, "failed")

    def test_case_page_shows_only_my_notifications_unless_edu(self):
        case = self.drafting_case()
        self.upload(case)
        case = self.act(case, "submit_proposal", self.student)
        for user, visible in ((self.advisor, True), (self.edu, True), (self.student, False)):
            self.client.force_login(user)
            page = self.client.get(reverse("case_detail", args=[case.pk]))
            (self.assertContains if visible else self.assertNotContains)(page, "اعلان‌های ایمیلی")


class ReminderTests(AgentTestCase):
    """«سفر در زمان» با پارامتر today؛ فقط مهلت «تحویل پیشنهاد» همین پرونده بررسی می‌شود."""

    def setUp(self):
        super().setUp()
        self.case = self.drafting_case()
        self.deadline = self.case.deadlines.get(kind=DeadlineKind.PROPOSAL_SUBMISSION)
        self.due = self.deadline.due_date

    def reminders(self, day):
        return [i for i in services.due_reminders(day) if i["key"].startswith(f"reminder:{self.deadline.pk}:")]

    def stages(self, day):
        """[(مرحله، نام کاربری گیرنده)] برای یادآوری‌های سررسیدشده در این روز."""
        return sorted((i["key"].split(":")[2], i["to"].split("@")[0]) for i in self.reminders(day))

    def test_timeline_follows_the_rules_exactly(self):
        expected = {20: None, 15: None, 14: "d14", 10: "d14", 7: "d7", 3: "d3", 1: "d1", 0: "d0",
                    -1: "overdue1", -8: "overdue2"}
        for days_left, stage in expected.items():
            day = self.due - timedelta(days=days_left)
            self.assertEqual(self.stages(day), [(stage, "student")] if stage else [], days_left)

    def test_each_stage_is_sent_once_per_recipient(self):
        day = self.due - timedelta(days=7)
        (item,) = self.reminders(day)
        self.assertIn("۷ روز مانده", item["subject"])
        self.assertEqual(item["facts"]["deadline"]["days_left"], 7)
        self.assertIn("ارسال پیشنهاد برای استاد راهنما", item["facts"]["your_actions"])
        self.assertIn("student گرامی،", item["fallback_body"])
        self.api("api_notification_record", data={"key": item["key"], "status": "sent", "subject": "s", "body": "b"})
        self.assertEqual(self.stages(day), [])
        self.assertEqual(self.stages(day + timedelta(days=1)), [])                     # هنوز همان مرحله‌ی d7
        self.assertEqual(self.stages(self.due - timedelta(days=3)), [("d3", "student")])   # مرحله‌ی بعد
        notification = Notification.objects.get()
        self.assertEqual((notification.kind, notification.stage, notification.case, notification.deadline),
                         ("reminder", "d7", self.case, self.deadline))

    def test_failed_send_is_retried(self):
        day = self.due - timedelta(days=7)
        (item,) = self.reminders(day)
        self.api("api_notification_record", data={"key": item["key"], "status": "failed"})
        self.assertEqual(self.stages(day), [("d7", "student")])

    def test_reminder_goes_to_whoever_must_act_plus_the_author(self):
        day = self.due - timedelta(days=3)
        wf.upload_document(self.case, self.student, DocKind.PROPOSAL, docx())
        case = self.do(self.case, "submit_proposal", self.student)
        self.assertEqual(self.stages(day), [("d3", "advisor"), ("d3", "student")])
        self.do(case, "advisor_approve", self.advisor)
        self.assertEqual(self.stages(day), [("d3", "edu"), ("d3", "student")])

    def test_satisfied_deadlines_lower_bounds_and_closed_cases_are_silent(self):
        def labels(day):
            return {i["facts"]["deadline"]["label"] for i in services.due_reminders(day)}

        wf.upload_document(self.case, self.student, DocKind.PROPOSAL, docx())
        case = self.do(self.case, "submit_proposal", self.student)
        case = self.do(case, "advisor_approve", self.advisor)
        case = self.do(case, "edu_register", self.edu)                  # مهلت تحویل پیشنهاد «انجام شد»
        self.assertEqual(self.stages(self.due), [])
        self.assertIn("پایان داوری در گروه آموزشی", labels(self.due))   # مهلت بعدی جای آن را گرفته است
        case = self.do(case, "assign_referees", self.head, referees=[self.ref1])
        case = self.do(case, "submit_review", self.ref1, decision="approve")
        case = self.do(case, "group_approve", self.head)
        case = self.do(case, "council_approve", self.edu)
        earliest = case.deadlines.get(kind=DeadlineKind.DEFENSE_EARLIEST).due_date
        for offset in (-3, 0, 5):                                       # «زودترین تاریخ دفاع» سررسید نیست
            self.assertNotIn("زودترین تاریخ مجاز دفاع", labels(earliest + timedelta(days=offset)))
        self.do(case, "cancel", self.edu, note="لغو")
        thesis_due = case.deadlines.get(kind=DeadlineKind.THESIS_DELIVERY).due_date
        self.assertEqual(services.due_reminders(thesis_due), [])        # پرونده‌ی بسته یادآوری ندارد

    def test_n8n_fetches_due_reminders_with_a_simulated_date(self):
        day = self.due - timedelta(days=1)
        data = self.api("api_reminders_due", method="get", query=f"?today={day.isoformat()}").json()
        self.assertEqual(data["count"], len(data["notifications"]))
        self.assertIn(f"reminder:{self.deadline.pk}:d1:{self.student.pk}", [n["key"] for n in data["notifications"]])
        today = self.api("api_reminders_due", method="get").json()
        self.assertNotIn(self.deadline.pk, [int(n["key"].split(":")[1]) for n in today["notifications"]])
        self.assertEqual(self.api("api_reminders_due", method="get", query="?today=tomorrow").status_code, 409)


# ───────────────────────── خاموش بودن عامل‌ها ─────────────────────────

@override_settings(AGENTS_ENABLED=False)
class DisabledTests(AgentTestCase):
    def test_portal_behaves_exactly_like_phase_two(self):
        case = self.drafting_case()
        self.upload(case)
        case = self.act(case, "submit_proposal", self.student)
        case = self.act(case, "advisor_approve", self.advisor)
        case = self.act(case, "edu_register", self.edu)
        self.n8n.assert_not_called()
        self.assertFalse(DocumentValidation.objects.exists() or Notification.objects.exists())
        self.client.force_login(self.head)
        page = self.client.get(reverse("case_action", args=[case.pk, "assign_referees"]))
        self.assertNotContains(page, "پیشنهاد سامانه")
        self.assertNotContains(self.client.get(reverse("case_detail", args=[case.pk])), "بررسی خودکار قالب")
        self.assertNotContains(self.client.get(reverse("dashboard")), "assistant-panel")
        self.assertEqual(self.client.post(reverse("assistant_ask"), "{}", content_type="application/json").status_code, 404)
